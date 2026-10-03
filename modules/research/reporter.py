"""
Excel report generator for KDP research results.
Produces color-coded, formatted reports ready to share or archive.
"""
import logging
from pathlib import Path
from datetime import datetime
from typing import List, Dict

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Color palette
C_DARK = "1A252F"
C_HEADER_BG = "2C3E50"
C_HEADER_FG = "FFFFFF"
C_ALT_ROW = "F2F6FA"
C_GREEN = "A8E6A3"
C_YELLOW = "FFF3B0"
C_ORANGE = "FFD580"
C_RED = "FFB3B3"
C_ACCENT = "3498DB"
C_WHITE = "FFFFFF"


def _score_color(score: float) -> str:
    if score >= 75:
        return C_GREEN
    elif score >= 60:
        return C_YELLOW
    elif score >= 45:
        return C_ORANGE
    else:
        return C_RED


def _reviews_color(reviews: int) -> str:
    if reviews <= 25:
        return C_GREEN
    elif reviews <= 100:
        return C_YELLOW
    elif reviews <= 300:
        return C_ORANGE
    else:
        return C_RED


def _fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _header_style(cell, size: int = 11):
    cell.fill = _fill(C_HEADER_BG)
    cell.font = Font(color=C_HEADER_FG, bold=True, size=size)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _thin_border() -> Border:
    thin = Side(style="thin", color="DDDDDD")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _set_col_width(ws, col: int, width: float):
    ws.column_dimensions[get_column_letter(col)].width = width


def _build_summary_sheet(ws, keyword: str, scores: Dict, products: List[Dict]):
    ws.sheet_view.showGridLines = False
    _set_col_width(ws, 1, 28)
    _set_col_width(ws, 2, 18)
    _set_col_width(ws, 3, 42)

    row = 1

    # Title band
    ws.merge_cells(f"A{row}:C{row}")
    cell = ws[f"A{row}"]
    cell.value = f"KDP Research Report"
    cell.fill = _fill(C_HEADER_BG)
    cell.font = Font(color=C_WHITE, bold=True, size=18)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 42

    row += 1
    ws.merge_cells(f"A{row}:C{row}")
    sub = ws[f"A{row}"]
    sub.value = f'Keyword: "{keyword}"  |  Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}'
    sub.fill = _fill(C_ACCENT)
    sub.font = Font(color=C_WHITE, size=10)
    sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 22

    row += 2

    # Big opportunity score
    ws.merge_cells(f"A{row}:C{row}")
    score_cell = ws[f"A{row}"]
    score_cell.value = (
        f"{scores['tier_emoji']} OPPORTUNITY SCORE: {scores['overall']}/100 — {scores['tier']}"
    )
    score_cell.fill = _fill(_score_color(scores["overall"]))
    score_cell.font = Font(bold=True, size=14, color=C_DARK)
    score_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 32

    row += 1
    ws.merge_cells(f"A{row}:C{row}")
    rec = ws[f"A{row}"]
    rec.value = scores["recommendation"]
    rec.font = Font(italic=True, size=10, color="444444")
    rec.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28

    row += 2

    # Score breakdown header
    for col, label in enumerate(["Score Component", "Score /100", "Interpretation"], 1):
        cell = ws.cell(row=row, column=col)
        cell.value = label
        _header_style(cell)
    ws.row_dimensions[row].height = 24
    row += 1

    breakdown = [
        ("Competition Score", scores["competition_score"], "How easy it is to rank (higher = less competition)"),
        ("Demand Score", scores["demand_score"], "Market activity and search demand"),
        ("Pricing Score", scores["pricing_score"], "Price range alignment with KDP royalty sweet spot"),
        ("OVERALL", scores["overall"], scores["recommendation"]),
    ]

    for metric, score, interp in breakdown:
        ws.cell(row=row, column=1).value = metric
        ws.cell(row=row, column=1).font = Font(bold=(metric == "OVERALL"), size=10)

        score_cell = ws.cell(row=row, column=2)
        score_cell.value = f"{score}/100"
        score_cell.fill = _fill(_score_color(score))
        score_cell.font = Font(bold=True, size=10, color=C_DARK)
        score_cell.alignment = Alignment(horizontal="center")

        ws.cell(row=row, column=3).value = interp
        ws.cell(row=row, column=3).font = Font(size=9, italic=True)
        ws.row_dimensions[row].height = 20
        row += 1

    row += 2

    # Key stats
    for col, label in enumerate(["Metric", "Value"], 1):
        cell = ws.cell(row=row, column=col)
        cell.value = label
        _header_style(cell)
    ws.row_dimensions[row].height = 24
    row += 1

    stats = [
        ("Products Analyzed", scores["total_products_analyzed"]),
        ("Average Reviews (top 10)", f"{int(scores['avg_reviews'])}"),
        ("Average Rating (top 10)", f"{scores['avg_rating']:.1f} / 5.0"),
        ("Average Price", f"${scores['avg_price']:.2f}" if scores["avg_price"] else "N/A"),
        ("Keyword Searched", keyword),
        ("Report Date", datetime.now().strftime("%Y-%m-%d")),
    ]

    for stat, val in stats:
        ws.cell(row=row, column=1).value = stat
        ws.cell(row=row, column=1).font = Font(size=10)
        ws.cell(row=row, column=2).value = val
        ws.cell(row=row, column=2).font = Font(bold=True, size=10)
        if row % 2 == 0:
            for col in [1, 2]:
                ws.cell(row=row, column=col).fill = _fill(C_ALT_ROW)
        row += 1


def _build_products_sheet(ws, products: List[Dict]):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    columns = [
        ("#", 5),
        ("ASIN", 12),
        ("Title", 48),
        ("Author", 22),
        ("Rating", 8),
        ("Reviews", 10),
        ("Price $", 9),
        ("Format", 13),
        ("BSR", 12),
        ("Competition", 13),
    ]

    for col, (header, width) in enumerate(columns, 1):
        _set_col_width(ws, col, width)
        cell = ws.cell(row=1, column=col)
        cell.value = header
        _header_style(cell)
    ws.row_dimensions[1].height = 26

    for i, p in enumerate(products, 2):
        reviews = p.get("reviews", 0)
        row_data = [
            p.get("position", i - 1),
            p.get("asin", ""),
            p.get("title", ""),
            p.get("author", ""),
            p.get("rating", 0),
            reviews,
            p.get("price", 0),
            p.get("format", ""),
            p.get("bsr_main", ""),
            "Low" if reviews <= 25 else "Medium" if reviews <= 100 else "High" if reviews <= 300 else "Very High",
        ]

        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col)
            cell.value = val
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical="center")
            if i % 2 == 0:
                cell.fill = _fill(C_ALT_ROW)

        # Color-code the reviews cell
        reviews_cell = ws.cell(row=i, column=6)
        reviews_cell.fill = _fill(_reviews_color(reviews))
        reviews_cell.font = Font(bold=True, size=9)

        # Color-code the competition label
        comp_cell = ws.cell(row=i, column=10)
        comp_cell.fill = _fill(_reviews_color(reviews))
        comp_cell.font = Font(bold=True, size=9)
        comp_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[i].height = 18


def _build_keywords_sheet(ws, seed: str, suggestions: List):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    _set_col_width(ws, 1, 45)
    _set_col_width(ws, 2, 20)
    _set_col_width(ws, 3, 10)

    for col, label in enumerate(["Keyword", "Source / Seed", "Level"], 1):
        cell = ws.cell(row=1, column=col)
        cell.value = label
        _header_style(cell)
    ws.row_dimensions[1].height = 26

    if suggestions and isinstance(suggestions[0], dict):
        rows = suggestions
    else:
        rows = [{"keyword": kw, "source": seed, "level": 1} for kw in suggestions]

    for i, item in enumerate(rows, 2):
        ws.cell(row=i, column=1).value = item.get("keyword", "")
        ws.cell(row=i, column=2).value = item.get("source", seed)
        ws.cell(row=i, column=3).value = item.get("level", 1)

        for col in [1, 2, 3]:
            cell = ws.cell(row=i, column=col)
            cell.font = Font(size=10)
            if i % 2 == 0:
                cell.fill = _fill(C_ALT_ROW)

        ws.row_dimensions[i].height = 18


def generate_research_report(
    keyword: str,
    products: List[Dict],
    scores: Dict,
    keyword_suggestions,
    output_path: str,
) -> str:
    """
    Generate a formatted Excel research report.
    Returns the path to the created file.
    """
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    _build_summary_sheet(ws_summary, keyword, scores, products)

    ws_products = wb.create_sheet("Products")
    _build_products_sheet(ws_products, products)

    ws_keywords = wb.create_sheet("Keywords")
    _build_keywords_sheet(ws_keywords, keyword, keyword_suggestions)

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    logger.info(f"Report saved: {path}")
    return str(path)
