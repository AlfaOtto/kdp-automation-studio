"""Amazon Ads (Sponsored Products) helper for the Analytics page.

Small-budget philosophy (the user typically has ~$60): be surgical — exact-match
long-tail keywords, low bids, kill losers fast, concentrate on 2-3 winners.

Provides: a per-book campaign suggestion (keywords + bids + negatives + settings),
the management rules, persistence of a tracking sheet, ACOS/decision computation,
and an Excel export.
"""
from __future__ import annotations
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

CAMPAIGN_FILE = "ads_campaign.json"

# ── The rules (shown in the UI) ───────────────────────────────────────────────
RULES: List[str] = [
    "Lance la pub SEULEMENT après ~5-10 avis (sinon budget gaspillé).",
    "Le compte doit être ACTIF (pas suspendu).",
    "Sponsored Products · Ciblage MANUEL · Correspondance EXACTE uniquement "
    "(jamais Auto/Broad avec un petit budget).",
    "Budget ~4 $/jour → les 60 $ durent ~15 jours. Enchères 0,25-0,35 $.",
    "Vérifie tous les 2-3 jours et remplis Dépense / Clics / Ventes.",
    "⛔ METS EN PAUSE un mot-clé qui dépense > 3 $ SANS vente.",
    "✅ GARDE (et +0,05 $) un mot-clé qui vend avec ACOS < 70 %.",
    "Concentre le budget restant sur les 2-3 mots-clés gagnants.",
    "Objectif = trouver des mots-clés rentables + booster le rang organique.",
]

DEFAULT_SETTINGS = {
    "type": "Sponsored Products",
    "targeting": "Manuel · Correspondance exacte",
    "daily_budget": 4.0,
    "default_bid": 0.30,
    "total_budget": 60.0,
}


def _clean(k: str) -> str:
    return re.sub(r"\s+", " ", str(k or "").strip().lower())[:50]


# Accessory/part words → block AMAZON shoppers who want the physical tool, not a book.
_ACC_COOK = ["machine", "oven", "basket", "liners", "accessories", "parts",
             "rack", "cover", "tray", "replacement", "seal"]
_ACC_TOOL = ["kit", "tools", "tool set", "supplies", "set", "starter kit",
             "machine", "equipment", "materials"]
# Other cooking METHODS a shopper might want instead (generic, no brand names).
_OTHER_METHODS = ["slow cooker", "pressure cooker", "microwave", "grill",
                  "sheet pan", "dutch oven"]
# Angle → non-book PRODUCT searches to exclude (supplements, snacks, gear…).
_ANGLE_NEG = {
    "protein":  ["protein powder", "protein shake", "protein bars", "protein supplement"],
    "keto":     ["keto snacks", "keto bars", "keto supplements"],
    "collagen": ["collagen powder", "collagen supplement"],
    "vegan":    ["vegan protein powder", "vegan snacks"],
    "gardening": ["seeds", "soil", "fertilizer", "planters", "garden tools"],
    "woodworking": ["wood", "lumber", "power tools", "hand tools"],
    "resin":    ["resin kit", "epoxy resin", "resin molds", "resin supplies"],
    "real estate": ["real estate license", "realtor course", "mls"],
}


def _rule_negatives(listing: Dict, core: str, anchor: str) -> List[str]:
    """AMAZON-appropriate negatives, derived from THIS book. Blocks shoppers who
    want the appliance/tool, accessories, supplements, or a different method —
    NOT web/piracy junk (free/pdf/youtube) which never applies on Amazon."""
    pt = listing.get("product_type", "")
    niche = (listing.get("niche", "") or "").lower()
    a = (anchor or core or "").strip()
    negs: List[str] = []
    if a:
        acc = _ACC_COOK if pt == "recipe_book" else _ACC_TOOL
        negs += [f"{a} {w}" for w in acc]        # e.g. "air fryer machine", "air fryer basket"
    if pt == "recipe_book":
        negs += _OTHER_METHODS
    for key, extra in _ANGLE_NEG.items():
        if key in niche:
            negs += extra
    return negs


def _core_anchor(niche: str) -> tuple:
    """Return (core, anchor). core = niche minus format/audience words;
    anchor = the distinctive subject phrase every ad keyword must contain."""
    core = re.sub(r"\b(for\s+beginners?|beginners?|guide|book|cookbook|handbook|"
                  r"journal|recipes?|edition|volume\s*\d*)\b", " ", niche or "",
                  flags=re.I)
    core = re.sub(r"\s+", " ", core).strip().lower()
    toks = [t for t in core.split() if t]
    anchor = " ".join(toks[-2:]) if len(toks) >= 2 else (toks[-1] if toks else core)
    return core, anchor


def suggest_campaign(listing: Dict) -> Dict:
    """Build a TIGHT, precise exact-match campaign. Every keyword is anchored on
    the subject (relevance guard) and carries a book-buying signal (format noun
    or audience) so the budget isn't wasted on tool/appliance shoppers."""
    niche = listing.get("niche", "") or ""
    pt = listing.get("product_type", "")
    own_kw = listing.get("keywords", []) or []
    core, anchor = _core_anchor(niche)
    core = core or niche.strip().lower()
    anchor = anchor or core
    typ = ("cookbook" if pt == "recipe_book"
           else "guide" if pt == "how_to" else "book")
    is_facts = bool(listing.get("_facts"))
    if is_facts:
        typ = "book"

    # High-intent templates. Build from BOTH the full subject (e.g. "high protein
    # air fryer") AND the short anchor (e.g. "air fryer") → more variety, more of
    # the phrases buyers actually type. Every keyword includes a book-buying signal.
    bases = [core]
    if anchor and anchor != core:
        bases.append(anchor)

    def _for_base(base):
        t = [f"{base} {typ}", f"best {base} {typ}", f"easy {base} {typ}",
             f"{base} {typ} for beginners", f"{base} {typ} 2026"]
        if pt == "recipe_book":
            t += [f"{base} recipe book", f"{base} recipes",
                  f"{base} meal prep cookbook", f"healthy {base} cookbook"]
        elif pt == "how_to":
            t += [f"how to {base}", f"{base} for beginners",
                  f"beginner {base} guide", f"{base} step by step"]
        else:
            t += [f"{base} for beginners", f"best {base} book"]
        return t

    tmpl = []
    for b in bases:
        tmpl += _for_base(b)

    # Also add the book's own keywords that are on-subject (contain the anchor).
    on_topic_own = [k for k in own_kw if anchor and anchor in _clean(k)]

    try:
        from modules.creation.content_generator import _has_brand
    except Exception:
        def _has_brand(_):  # pragma: no cover
            return False

    seen: List[str] = []
    out: List[str] = []
    for k in tmpl + on_topic_own:
        kk = _clean(k)
        # relevance guard: contains the subject anchor; 2-8 words; no brand
        if (kk and kk not in seen and 2 <= len(kk.split()) <= 8
                and (anchor in kk or anchor == core) and not _has_brand(kk)):
            seen.append(kk)
            out.append(kk)
    out = out[:18]
    bid = DEFAULT_SETTINGS["default_bid"]
    keywords = [{"keyword": k, "bid": (0.35 if i < 2 else bid)} for i, k in enumerate(out)]

    # AMAZON-appropriate negatives (appliance / accessories / supplements / other
    # methods) derived from THIS book. Guard: never a phrase contained in one of
    # our own target keywords (would block our ads).
    kw_texts = [k["keyword"] for k in keywords]
    negs = _rule_negatives(listing, core, anchor)
    negatives, nseen = [], set()
    for n in negs:
        n = _clean(n)
        if n and n not in nseen and not any(n in kt for kt in kw_texts):
            nseen.add(n)
            negatives.append(n)
    return {
        "settings": dict(DEFAULT_SETTINGS),
        "keywords": keywords,
        "negatives": negatives,
    }


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(s or "").lower()).strip("-")[:60] or "book"


def _claude_negatives(listing: Dict, anthropic_key: str, s=lambda m: None) -> List[str]:
    """One cheap Haiku call → intelligent, Amazon-specific negative keywords."""
    from anthropic import Anthropic
    client = Anthropic(api_key=anthropic_key)
    title = listing.get("title", "")
    sub = listing.get("subtitle", "")
    niche = listing.get("niche", "")
    pt = listing.get("product_type", "")
    typ = ("cookbook" if pt == "recipe_book"
           else "how-to / guide book" if pt == "how_to" else "book")
    s("🤖 Génération de négatifs Amazon intelligents (Claude)…")
    prompt = f"""You are an Amazon Ads (Sponsored Products) expert for KDP books.
Book: "{title}: {sub}" — a {typ} about "{niche}".

List 15-22 NEGATIVE keywords for an AMAZON Sponsored Products campaign for THIS book.
Negatives = searches ON AMAZON where the shopper wants something OTHER than this book:
- the physical appliance / tool / device itself, and its accessories / parts
- supplements, snacks, gear, or other physical PRODUCTS related to the topic
- a DIFFERENT method / appliance / diet / format than this book covers
- clearly different sub-topics that would waste ad budget

STRICT RULES:
- Amazon context ONLY. NEVER include web/piracy terms (free, pdf, epub, torrent,
  youtube, app, download, audiobook, kindle unlimited) — irrelevant on Amazon.
- NEVER include any brand or trademark name.
- NEVER include a phrase a buyer of THIS book would search (don't block your own sales).
- 1-4 words each, lowercase, no duplicates.
Return ONLY a JSON array of strings."""
    msg = client.messages.create(
        model="claude-haiku-4-5", max_tokens=700,
        messages=[{"role": "user", "content": prompt}])
    txt = "".join(getattr(b, "text", "") for b in msg.content
                  if getattr(b, "type", "") == "text")
    m = re.search(r"\[.*\]", txt, re.S)
    arr = json.loads(m.group(0)) if m else []
    return [str(x).strip().lower() for x in arr if str(x).strip()]


def smart_negatives(listing: Dict, anthropic_key: str = "", cache_dir=None,
                    refresh: bool = False, status_cb=None) -> tuple:
    """Intelligent Amazon negatives via Claude, CACHED per niche (0 $ after the
    first generation). Falls back to rule-based negatives. Returns (negatives, source)."""
    def s(m):
        if status_cb:
            status_cb(m)
    slug = _slug(listing.get("niche") or listing.get("title"))
    cache_file = (Path(cache_dir) / "ad_negatives" / f"{slug}.json") if cache_dir else None
    if cache_file and cache_file.exists() and not refresh:
        try:
            return json.loads(cache_file.read_text(encoding="utf-8")), "cache (0 $)"
        except Exception:
            pass
    negs, src = [], "IA (Claude)"
    if anthropic_key:
        try:
            negs = _claude_negatives(listing, anthropic_key, s)
        except Exception as e:
            s(f"IA négatifs indisponible : {str(e)[:60]}")
    if not negs:
        core, anchor = _core_anchor(listing.get("niche", "") or "")
        negs = _rule_negatives(listing, core, anchor)
        src = "règles"
    core, anchor = _core_anchor(listing.get("niche", "") or "")
    seen, out = set(), []
    for n in negs:
        n = _clean(n)
        if n and n not in seen and n not in (core, anchor):
            seen.add(n)
            out.append(n)
    out = out[:30]
    if cache_file and out:
        cache_file.parent.mkdir(parents=True, exist_ok=True)
        cache_file.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    return out, src


def extract_competitor_asins(keywords: List[str], max_per_kw: int = 15,
                             exclude_sponsored: bool = True,
                             status_cb=None) -> List[Dict]:
    """Scrape Amazon for the given keyword(s) and return the most in-demand
    competitor books (ASIN + title + reviews + rating + price), sorted by number
    of reviews (a good proxy for demand). Use these ASINs in an Amazon Ads
    Product Targeting campaign."""
    def s(m):
        if status_cb:
            status_cb(m)
    try:
        from modules.research.scraper import AmazonScraper
    except Exception as e:
        s(f"Scraper indisponible : {e}")
        return []
    seen: Dict[str, Dict] = {}
    try:
        with AmazonScraper() as sc:
            for kw in keywords:
                kw = str(kw or "").strip()
                if not kw:
                    continue
                s(f"🔍 Recherche Amazon : « {kw} »")
                try:
                    prods = sc.search_books(keyword=kw, max_results=max_per_kw)
                except Exception as e:
                    s(f"⚠️ « {kw} » : {str(e)[:60]}")
                    prods = []
                for p in prods:
                    a = (p.get("asin") or "").strip()
                    if not a:
                        continue
                    if exclude_sponsored and p.get("sponsored"):
                        continue
                    rev = int(p.get("reviews") or 0)
                    if a not in seen or rev > (seen[a].get("reviews") or 0):
                        seen[a] = {
                            "asin": a,
                            "title": (p.get("title") or "")[:90],
                            "reviews": rev,
                            "rating": p.get("rating"),
                            "price": p.get("price"),
                            "keyword": kw,
                            "url": p.get("url") or f"https://www.amazon.com/dp/{a}",
                        }
    except Exception as e:
        s(f"Extraction interrompue : {str(e)[:80]}")
    return sorted(seen.values(), key=lambda x: x.get("reviews") or 0, reverse=True)


ASINS_FILE = "competitor_asins.json"


def load_competitor_asins(product_dir: Path) -> List[Dict]:
    fp = Path(product_dir) / ASINS_FILE
    if fp.exists():
        try:
            return json.loads(fp.read_text(encoding="utf-8"))
        except Exception:
            return []
    return []


def save_competitor_asins(product_dir: Path, asins: List[Dict]) -> None:
    (Path(product_dir) / ASINS_FILE).write_text(
        json.dumps(asins, ensure_ascii=False, indent=2), encoding="utf-8")


def asins_from_pool(learning_dir: Path, listing: Dict, limit: int = 25) -> List[Dict]:
    """Reuse competitor books ALREADY scraped (learning/competitor_samples.json)
    that match this book's topic → 0 new scraping, 0 cost."""
    fp = Path(learning_dir) / "competitor_samples.json"
    if not fp.exists():
        return []
    try:
        pool = json.loads(fp.read_text(encoding="utf-8"))
    except Exception:
        return []
    toks = set(re.findall(r"[a-z]{4,}", (listing.get("niche", "") or "").lower()))
    for k in listing.get("keywords", []) or []:
        toks |= set(re.findall(r"[a-z]{4,}", str(k).lower()))
    toks -= {"beginners", "beginner", "cookbook", "guide", "book", "recipes",
             "recipe", "journal", "edition"}
    items = pool.values() if isinstance(pool, dict) else pool
    out = []
    for it in items:
        if not isinstance(it, dict):
            continue
        asin = it.get("asin")
        title = (it.get("title") or "").lower()
        cats = " ".join((it.get("analysis", {}) or {}).get("topic_categories", [])).lower()
        hay = title + " " + cats
        score = sum(1 for t in toks if t in hay)
        if asin and score >= 1:
            out.append({"asin": asin, "title": (it.get("title") or "")[:90],
                        "reviews": it.get("reviews", 0), "rating": it.get("rating"),
                        "price": it.get("price"), "keyword": "(déjà scrapé)",
                        "_score": score})
    out.sort(key=lambda x: (x.get("_score", 0), x.get("reviews") or 0), reverse=True)
    for o in out:
        o.pop("_score", None)
    return out[:limit]


def get_competitor_asins(product_dir: Path, keywords: List[str], listing: Dict,
                         learning_dir: Path, refresh: bool = False,
                         status_cb=None) -> tuple:
    """Return (asins, source). By default REUSES what's already scraped (per-book
    cache, then the global pool) — NO new scraping. Only scrapes Amazon live when
    refresh=True, then caches the result for next time."""
    if not refresh:
        cache = load_competitor_asins(product_dir)
        if cache:
            return cache, "cache livre (0 scraping)"
        pool = asins_from_pool(learning_dir, listing)
        if pool:
            return pool, "pool déjà scrapé (0 scraping)"
        return [], "aucun cache — clique « Rafraîchir » pour scraper une fois"
    # refresh → scrape live once, merge with pool, cache
    live = extract_competitor_asins(keywords, status_cb=status_cb)
    merged = {a["asin"]: a for a in asins_from_pool(learning_dir, listing)}
    for a in live:
        merged[a["asin"]] = a          # live data (with reviews) wins
    res = sorted(merged.values(), key=lambda x: x.get("reviews") or 0, reverse=True)
    if res:
        save_competitor_asins(product_dir, res)
    return res, "scrapé + mis en cache"


def default_royalty(listing: Dict) -> float:
    """A sensible default royalty/sale for the ACOS math (ebook-oriented)."""
    pt = listing.get("product_type", "")
    # illustrated ebook (70% minus delivery) ≈ 2.6-3.5; guides a bit higher
    return 3.0 if pt == "recipe_book" else 3.2


# ── Tracking persistence ──────────────────────────────────────────────────────
def load_campaign(product_dir: Path) -> Optional[Dict]:
    fp = Path(product_dir) / CAMPAIGN_FILE
    if fp.exists():
        try:
            return json.loads(fp.read_text(encoding="utf-8"))
        except Exception:
            return None
    return None


HISTORY_FILE = "ads_history.jsonl"


def append_history(product_dir: Path, campaign: Dict) -> bool:
    """Append a timestamped snapshot of the campaign to the per-book history
    (append-only JSONL). Skips if identical to the last snapshot. Returns True
    if a new snapshot was written."""
    snap = {
        "ts": datetime.now().isoformat(timespec="seconds"),
        "royalty": campaign.get("royalty"),
        "budget": campaign.get("budget"),
        "totals": totals(campaign),
        "rows": [dict(r) for r in campaign.get("rows", [])],
    }
    fp = Path(product_dir) / HISTORY_FILE
    try:
        if fp.exists():
            lines = [l for l in fp.read_text(encoding="utf-8").splitlines() if l.strip()]
            if lines:
                prev = json.loads(lines[-1])
                if (prev.get("rows") == snap["rows"]
                        and prev.get("totals", {}).get("spend") == snap["totals"]["spend"]):
                    return False   # nothing changed → don't duplicate
    except Exception:
        pass
    with open(fp, "a", encoding="utf-8") as f:
        f.write(json.dumps(snap, ensure_ascii=False) + "\n")
    return True


def load_history(product_dir: Path) -> List[Dict]:
    """Return all saved snapshots (oldest → newest) for this book."""
    fp = Path(product_dir) / HISTORY_FILE
    out: List[Dict] = []
    if fp.exists():
        for line in fp.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line:
                try:
                    out.append(json.loads(line))
                except Exception:
                    pass
    return out


def keyword_history(product_dir: Path, keyword: str) -> List[Dict]:
    """Evolution of one keyword across snapshots: [{ts, spend, clicks, orders}]."""
    kw = re.sub(r"\s+", " ", str(keyword or "").strip().lower())
    out = []
    for snap in load_history(product_dir):
        for r in snap.get("rows", []):
            if r.get("keyword") == kw:
                out.append({"ts": snap["ts"], "spend": r.get("spend", 0),
                            "clicks": r.get("clicks", 0), "orders": r.get("orders", 0)})
                break
    return out


def save_campaign(product_dir: Path, campaign: Dict) -> None:
    campaign["updated_at"] = datetime.now().isoformat(timespec="seconds")
    (Path(product_dir) / CAMPAIGN_FILE).write_text(
        json.dumps(campaign, ensure_ascii=False, indent=2), encoding="utf-8")
    # Always keep a dated snapshot in the per-book history.
    append_history(product_dir, campaign)


def new_campaign(listing: Dict) -> Dict:
    sug = suggest_campaign(listing)
    roy = default_royalty(listing)
    rows = [{"keyword": k["keyword"], "bid": k["bid"],
             "spend": 0.0, "clicks": 0, "orders": 0} for k in sug["keywords"]]
    return {
        "settings": sug["settings"],
        "negatives": sug["negatives"],
        "royalty": roy,
        "budget": sug["settings"]["total_budget"],
        "rows": rows,
    }


# ── ACOS + decision ───────────────────────────────────────────────────────────
def compute_row(row: Dict, royalty: float) -> Dict:
    spend = float(row.get("spend") or 0)
    clicks = int(row.get("clicks") or 0)
    orders = int(row.get("orders") or 0)
    revenue = round(orders * float(royalty or 0), 2)
    acos = round(spend / revenue, 2) if revenue > 0 else None
    if spend == 0 and clicks == 0:
        decision = "—"
    elif orders == 0 and spend > 3:
        decision = "⛔ PAUSE"
    elif orders > 0 and acos is not None and acos < 0.70:
        decision = "✅ Garder (+0,05 $)"
    elif orders > 0:
        decision = "⚠️ ACOS élevé — baisser l'enchère"
    else:
        decision = "⏳ Laisser tourner"
    return {"revenue": revenue, "acos": acos, "decision": decision}


def totals(campaign: Dict) -> Dict:
    roy = float(campaign.get("royalty") or 0)
    spend = sum(float(r.get("spend") or 0) for r in campaign.get("rows", []))
    orders = sum(int(r.get("orders") or 0) for r in campaign.get("rows", []))
    revenue = round(orders * roy, 2)
    acos = round(spend / revenue, 2) if revenue > 0 else None
    return {
        "spend": round(spend, 2),
        "orders": orders,
        "revenue": revenue,
        "acos": acos,
        "remaining": round(float(campaign.get("budget") or 0) - spend, 2),
    }


def _num(v) -> float:
    """Parse a number that may carry currency symbols, %, commas, spaces."""
    try:
        s = re.sub(r"[^0-9.\-]", "", str(v).replace(",", "."))
        # collapse accidental double dots from thousand separators
        if s.count(".") > 1:
            head, _, tail = s.rpartition(".")
            s = head.replace(".", "") + "." + tail
        return float(s) if s not in ("", "-", ".") else 0.0
    except Exception:
        return 0.0


def parse_ads_report(path_or_buffer, filename: str = "") -> tuple:
    """Read an Amazon Ads export (.xlsx/.csv) and aggregate Spend/Clicks/Orders
    per keyword (or targeting / search term). Returns (report_dict, error_or_None)
    where report_dict = {keyword_lower: {spend, clicks, orders}}."""
    try:
        import pandas as pd
    except Exception:
        return {}, "pandas indisponible"
    name = (filename or getattr(path_or_buffer, "name", "") or str(path_or_buffer)).lower()
    try:
        if name.endswith(".csv"):
            df = pd.read_csv(path_or_buffer)
        else:
            df = pd.read_excel(path_or_buffer)     # first sheet
    except Exception as e:
        return {}, f"Lecture impossible : {str(e)[:80]}"
    if df is None or df.empty:
        return {}, "Fichier vide"

    cols = {c: str(c).strip().lower() for c in df.columns}

    def find(subs, exclude=()):
        for c, low in cols.items():
            if any(s in low for s in subs) and not any(e in low for e in exclude):
                return c
        return None

    kw_col = (find(["keyword"], exclude=["match", "id", "bid", "state", "type"])
              or find(["targeting"], exclude=["id", "type"])
              or find(["customer search term", "search term"]))
    spend_col = find(["spend", "cost"])
    clicks_col = find(["clicks"])
    orders_col = (find(["7 day total orders", "14 day total orders", "total orders"])
                  or find(["orders"])
                  or find(["purchases"]))
    if kw_col is None:
        return {}, ("Colonne mot-clé introuvable — exporte le rapport « Search term » "
                    "ou « Targeting » depuis Amazon Ads.")

    agg: Dict[str, Dict] = {}
    for _, r in df.iterrows():
        k = re.sub(r"\s+", " ", str(r.get(kw_col, "")).strip().lower())
        if not k or k in ("nan", "total", "—", "-"):
            continue
        d = agg.setdefault(k, {"spend": 0.0, "clicks": 0, "orders": 0})
        if spend_col is not None:
            d["spend"] += _num(r.get(spend_col))
        if clicks_col is not None:
            d["clicks"] += int(_num(r.get(clicks_col)))
        if orders_col is not None:
            d["orders"] += int(_num(r.get(orders_col)))
    return agg, None


def merge_report(campaign: Dict, report: Dict,
                 add_new: bool = True, max_rows: int = 40) -> tuple:
    """Fill the campaign rows from a parsed report (match by keyword). Optionally
    add new keywords that got clicks (discovery). Returns (n_updated, n_added)."""
    idx = {r["keyword"]: r for r in campaign.get("rows", [])}
    updated = added = 0
    default_bid = campaign.get("settings", {}).get("default_bid", 0.30)
    for k, d in report.items():
        if k in idx:
            idx[k]["spend"] = round(d["spend"], 2)
            idx[k]["clicks"] = int(d["clicks"])
            idx[k]["orders"] = int(d["orders"])
            updated += 1
        elif add_new and d["clicks"] > 0 and len(campaign["rows"]) < max_rows:
            campaign["rows"].append({
                "keyword": k, "bid": default_bid,
                "spend": round(d["spend"], 2), "clicks": int(d["clicks"]),
                "orders": int(d["orders"])})
            added += 1
    return updated, added


def export_xlsx(product_dir: Path, campaign: Dict, slug: str) -> Optional[str]:
    """Write the tracking sheet as an .xlsx next to the book. Returns path or None."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except Exception:
        return None
    roy = float(campaign.get("royalty") or 0)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Suivi Ads"
    ws.append([f"CAMPAGNE {slug}  ·  budget {campaign.get('budget')}$  ·  "
               f"royalty/vente {roy}$  ·  {campaign['settings']['targeting']}"])
    ws.append([])
    ws.append(["Mot-clé (Exact)", "Enchère ($)", "Dépense ($)", "Clics", "Ventes",
               "Revenu ($)", "ACOS", "Décision"])
    start = 4
    rows = campaign.get("rows", [])
    for i, r in enumerate(rows):
        rr = start + i
        ws.append([r["keyword"], r.get("bid", 0.3), r.get("spend", 0),
                   r.get("clicks", 0), r.get("orders", 0),
                   f"=E{rr}*{roy}", f'=IF(F{rr}=0,"",C{rr}/F{rr})',
                   f'=IF(C{rr}=0,"—",IF(AND(C{rr}>3,E{rr}=0),"PAUSE",'
                   f'IF(AND(E{rr}>0,G{rr}<0.7),"Garder","Surveiller")))'])
    tr = start + len(rows) + 1
    ws.append([])
    ws.append(["TOTAL", "", f"=SUM(C{start}:C{start+len(rows)-1})",
               f"=SUM(D{start}:D{start+len(rows)-1})",
               f"=SUM(E{start}:E{start+len(rows)-1})",
               f"=SUM(F{start}:F{start+len(rows)-1})", "", ""])
    ws.append([f"Budget restant : =", "", f"={campaign.get('budget')}-C{tr}"])
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    for c in range(1, 9):
        cell = ws.cell(row=3, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate([34, 12, 12, 8, 8, 11, 9, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws2 = wb.create_sheet("Règles")
    ws2.append(["RÈGLES"])
    for r in RULES:
        ws2.append(["• " + r])
    ws2.column_dimensions["A"].width = 95
    out = Path(product_dir) / f"ads_{slug}.xlsx"
    wb.save(out)
    return str(out)
